import base64
from typing import Optional
from datetime import datetime
from io import BytesIO
import urllib.parse

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel, EmailStr

from configs.database import get_db
from models.users import Users
from configs.auth import AuthService

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

router = APIRouter(prefix="/users", tags=["Users"])


class UsersOut(BaseModel):
    id: int
    name: Optional[str]
    user_name: Optional[str]
    user_phone: Optional[str]
    email: Optional[str]
    status: Optional[str]
    avatar: Optional[str]
    created_at: datetime

    model_config = {"from_attributes": True}


class UsersCreate(BaseModel):
    name: str
    user_name: str
    email: EmailStr
    password: str
    user_phone: Optional[str] = None
    status: Optional[str] = "active"


class UsersUpdate(BaseModel):
    name: Optional[str] = None
    user_name: Optional[str] = None
    user_phone: Optional[str] = None
    email: Optional[EmailStr] = None
    password: Optional[str] = None
    status: Optional[str] = None
    avatar: Optional[str] = None


def _get_users_query(db: Session):
    return db.query(Users).filter(Users.deleted_at == None).order_by(Users.id.desc())


def _attachment_header(filename: str) -> str:
    encoded = urllib.parse.quote(filename)
    return f'attachment; filename="{filename}"; filename*=UTF-8\'\'{encoded}'


def set_cell_border(cell, border_color: str = "CCCCCC") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for side in ("top", "left", "bottom", "right"):
        border_el = OxmlElement(f"w:{side}")
        border_el.set(qn("w:val"), "single")
        border_el.set(qn("w:sz"), "4")
        border_el.set(qn("w:color"), border_color)
        tcPr.append(border_el)


@router.get("/", response_model=list[UsersOut], summary="รายการ users")
def list_users(
    page: Optional[int] = Query(None, ge=1),
    limit: Optional[int] = Query(None, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    query = _get_users_query(db)
    if page is not None and limit is not None:
        query = query.offset((page - 1) * limit).limit(limit)
    return query.all()


@router.get("/me", response_model=UsersOut, summary="ข้อมูล user ของตัวเอง")
def get_me(current_user: Users = Depends(AuthService.get_current_user)):
    return current_user


@router.get("/export/pdf", summary="Export users เป็น PDF")
def export_users_pdf(
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    users = _get_users_query(db).all()

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontSize=16, spaceAfter=12)

    story = [
        Paragraph("Users Report", title_style),
        Paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 12),
    ]

    headers = ["ID", "Name", "Username", "Phone", "Email", "Status", "Created At"]
    data = [headers]
    for u in users:
        data.append([
            str(u.id), u.name or "-", u.user_name or "-",
            u.user_phone or "-", u.email or "-", u.status or "-",
            u.created_at.strftime("%Y-%m-%d") if u.created_at else "-",
        ])

    col_widths = [25*mm, 45*mm, 40*mm, 35*mm, 65*mm, 25*mm, 35*mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4FB")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    story.append(table)
    doc.build(story)
    buffer.seek(0)

    filename = f"users_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf"
    return {
        "filename": filename,
        "content_type": "application/pdf",
        "data": base64.b64encode(buffer.read()).decode(),
    }

@router.get("/export/excel", summary="Export users เป็น Excel")
def export_users_excel(
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    users = _get_users_query(db).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Users"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2E75B6")
    center_align = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["ID", "Name", "Username", "Phone", "Email", "Status", "Created At"]
    col_widths = [8, 25, 22, 18, 35, 12, 18]

    for col, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 22

    for row_idx, u in enumerate(users, start=2):
        row_data = [
            u.id, u.name or "", u.user_name or "", u.user_phone or "",
            u.email or "", u.status or "",
            u.created_at.strftime("%Y-%m-%d %H:%M") if u.created_at else "",
        ]
        fill = PatternFill("solid", fgColor="EEF4FB" if row_idx % 2 == 0 else "FFFFFF")
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                horizontal="center" if col in (1, 6, 7) else "left",
                vertical="center",
            )
            cell.border = border
            cell.fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"users_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return {
        "filename": filename,
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "data": base64.b64encode(buffer.read()).decode(),
    }


@router.get("/export/word", summary="Export users เป็น Word")
def export_users_word(
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    users = _get_users_query(db).all()

    doc = Document()
    section = doc.sections[0]
    section.orientation = 1
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.5)

    title = doc.add_heading("Users Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} | Total: {len(users)} records"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths_in = [0.4, 1.8, 1.6, 1.4, 2.5, 1.0, 1.4]
    for i, width in enumerate(col_widths_in):
        table.columns[i].width = Inches(width)

    headers = ["ID", "Name", "Username", "Phone", "Email", "Status", "Created At"]
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(11)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "2E75B6")
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        set_cell_border(hdr_cells[i])

    for idx, u in enumerate(users):
        row_cells = table.add_row().cells
        row_data = [
            str(u.id), u.name or "-", u.user_name or "-",
            u.user_phone or "-", u.email or "-", u.status or "-",
            u.created_at.strftime("%Y-%m-%d") if u.created_at else "-",
        ]
        fill_color = "EEF4FB" if (idx + 1) % 2 == 0 else "FFFFFF"
        for i, text in enumerate(row_data):
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 5, 6) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(10)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), fill_color)
            row_cells[i]._tc.get_or_add_tcPr().append(shading)
            set_cell_border(row_cells[i])

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"users_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.docx"
    return {
        "filename": filename,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "data": base64.b64encode(buffer.read()).decode(),
    }

# ====================== CRUD ======================
@router.get("/{user_id}", response_model=UsersOut, summary="ดู user ตาม ID")
def get_user(
    user_id: int,
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    user = db.query(Users).filter(Users.id == user_id, Users.deleted_at == None).first()
    if not user:
        raise HTTPException(status_code=404, detail=f"ไม่พบ user id={user_id}")
    return user


@router.post("/", response_model=UsersOut, status_code=status.HTTP_201_CREATED, summary="สร้าง user ใหม่")
def create_user(
    body: UsersCreate,
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    try:
        if db.query(Users).filter(Users.email == body.email).first():
            raise HTTPException(status_code=400, detail="Email นี้ถูกใช้งานแล้ว")

        user = Users(
            name=body.name,
            user_name=body.user_name,
            email=body.email,
            password=AuthService.hash_password(body.password),
            user_phone=body.user_phone,
            status=body.status,
            created_by=current_user.id,
            created_by_name=current_user.name,
        )
        db.add(user)
        db.commit()
        db.refresh(user)
        return user
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.put("/{user_id}", response_model=UsersOut, summary="อัปเดต user (PUT)")
def replace_user(
    user_id: int,
    body: UsersUpdate,
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    try:
        user = db.query(Users).filter(Users.id == user_id, Users.deleted_at == None).first()
        if not user:
            raise HTTPException(status_code=404, detail=f"ไม่พบ user id={user_id}")

        if body.name is not None:       user.name = body.name
        if body.user_name is not None:  user.user_name = body.user_name
        if body.user_phone is not None: user.user_phone = body.user_phone
        if body.email is not None:      user.email = body.email
        if body.password is not None:   user.password = AuthService.hash_password(body.password)
        if body.status is not None:     user.status = body.status
        if body.avatar is not None:     user.avatar = body.avatar

        user.updated_by = current_user.id
        user.updated_by_name = current_user.name

        db.commit()
        db.refresh(user)
        return user
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.patch("/{user_id}", response_model=UsersOut, summary="อัปเดตบางฟิลด์ (PATCH)")
def update_user(
    user_id: int,
    body: UsersUpdate,
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    return replace_user(user_id, body, db, current_user)


@router.delete("/{user_id}", status_code=status.HTTP_200_OK, summary="ลบ user (hard delete)")
def delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    current_user: Users = Depends(AuthService.get_current_user),
):
    try:
        user = db.query(Users).filter(Users.id == user_id).first()
        if not user:
            raise HTTPException(status_code=404, detail=f"ไม่พบ user id={user_id}")
        if user.id == current_user.id:
            raise HTTPException(status_code=400, detail="ไม่สามารถลบ user ของตัวเองได้")

        db.delete(user)
        db.commit()
        return {"message": f"ลบ user id={user_id} สำเร็จ"}
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))